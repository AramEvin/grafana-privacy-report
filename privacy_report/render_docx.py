"""Render findings.json into a DPIA-annex style DOCX using the Grafana Labs template look.

Pure python-docx: no Word install needed, runs in CI. Every visual block from the
template (cover banner, INFO/TIP/WARN callouts, numbered step blocks, code blocks,
troubleshooting table) is a small helper below so contributors can reuse them.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import remediation as rem
from .detectors import BY_KEY

# ---- brand palette ----------------------------------------------------------
ORANGE = "F46800"
YELLOW = "FBCA0A"
DARK = "181B1F"
INK = "24292E"
GREY = "6E7681"
LIGHT = "F4F5F5"
LINE = "D8DBDF"
SEV_FILL = {"Critical": "C4162A", "High": "F46800", "Medium": "E0B400", "Low": "56A64B"}
CALLOUT = {"INFO": "3D71D9", "TIP": "1A7F4B", "WARN": "C4162A"}
FONT = "Calibri"
MONO = "Consolas"
CONTENT_W = 17.0  # cm (A4 minus 2 cm margins)


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---- low-level OOXML helpers -----------------------------------------------

def shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (size_eighths, hex) or None for nil."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        b.append(el)


def table_borders(table, color: str | None = LINE, inside: bool = True):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if color is None or (not inside and edge.startswith("inside")):
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def table_margins(table, top=80, bottom=80, left=110, right=110):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def col_widths(table, widths_cm: list[float]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblW.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.append(c)


def para_border(paragraph, edge="bottom", size=8, color=ORANGE, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def add_field(paragraph, instr: str, placeholder: str = "", size=None, color=None):
    def run_with(el):
        r = paragraph.add_run()
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color
        r._r.append(el)
        return r

    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); run_with(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run_with(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); run_with(s)
    r = paragraph.add_run(placeholder)
    if size:
        r.font.size = size
    if color:
        r.font.color.rgb = color
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run_with(e)


def run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = rgb(color)
    if font:
        r.font.name = font
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    return r


def cell_text(cell, text, bold=False, size=9.5, color=INK, align=None, italic=False, font=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    run(p, str(text), bold=bold, size=size, color=color, italic=italic, font=font)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


# ---- document-level building blocks ----------------------------------------

class Report:
    def __init__(self, data: dict):
        self.data = data
        self.meta = data["meta"]
        self.doc = Document()
        self._setup()

    # styles, page, header/footer
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(1), Cm(1)

        st = d.styles
        normal = st["Normal"]
        normal.font.name = FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = rgb(INK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12

        for name, size, color, before, after in (("Heading 1", 18, DARK, 6, 10), ("Heading 2", 13, ORANGE, 14, 6),
                                                 ("Heading 3", 11, DARK, 10, 4)):
            s = st[name]
            s.font.name = FONT
            rf = s.element.rPr.rFonts
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(attr), FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(attr)) is not None:
                    del rf.attrib[qn(attr)]
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = rgb(color)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True

        sec.different_first_page_header_footer = True
        title = self.meta.get("title", "Telemetry Privacy Report")

        # the default Header/Footer styles carry center (4680) and right (9360) tabs; replace them
        for sname in ("Header", "Footer"):
            tabs = st[sname].element.pPr.find(qn("w:tabs"))
            if tabs is not None:
                st[sname].element.pPr.remove(tabs)

        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
        run(hp, "Grafana Labs", bold=True, size=9, color=DARK)
        run(hp, ".", bold=True, size=9, color=ORANGE)
        run(hp, "\t" + title, italic=True, size=9, color=GREY)
        para_border(hp, "bottom", 6, ORANGE, 4)

        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
        para_border(fp, "top", 4, LINE, 4)
        run(fp, "https://grafana.com   ·   ", size=8, color=GREY)
        run(fp, self.meta.get("classification", "CONFIDENTIAL — contains masked personal data findings"),
            bold=True, size=8, color=SEV_FILL["Critical"])
        run(fp, "\tPage ", size=8, color=GREY)
        add_field(fp, "PAGE", "1", Pt(8), rgb(GREY))
        run(fp, " of ", size=8, color=GREY)
        add_field(fp, "NUMPAGES", "1", Pt(8), rgb(GREY))

        # ask Word to refresh TOC / page numbers on open
        settings = d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    def p(self, text="", italic=False, color=None, size=None, after=6, bold=False, align=None):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(after)
        if align:
            para.alignment = align
        if text:
            run(para, text, italic=italic, color=color, size=size, bold=bold)
        return para

    def h1(self, text):
        h = self.doc.add_heading(level=1)
        num, _, rest = text.partition(" ")
        run(h, num + " ", color=ORANGE)
        run(h, rest)
        para_border(h, "bottom", 6, LINE, 4)
        return h

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def bullet(self, text, bold_lead: str | None = None):
        para = self.doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if bold_lead:
            run(para, bold_lead, bold=True)
        run(para, text)
        return para

    def spacer(self, pts=4):
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.line_spacing = Pt(pts)
        return sp

    # --- template blocks ------------------------------------------------------
    def callout(self, kind: str, text: str, lead: str | None = None):
        t = self.doc.add_table(rows=1, cols=2)
        table_borders(t, None)
        table_margins(t, 90, 90, 140, 140)
        col_widths(t, [1.9, CONTENT_W - 1.9])
        a, b = t.rows[0].cells
        shade(a, CALLOUT[kind])
        cell_text(a, kind, bold=True, size=9, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(b, LIGHT)
        cell_borders(b, left=(18, CALLOUT[kind]))
        p = b.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if lead:
            run(p, lead + " ", bold=True, size=9.5)
        run(p, text, size=9.5)
        b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        keep_row_together(t.rows[0])
        self.spacer(6)

    def code(self, text: str, label: str | None = None):
        if label:
            lp = self.p(label, bold=True, size=8.5, color=GREY, after=2)
            lp.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=1)
        table_borders(t, None)
        table_margins(t, 100, 100, 160, 120)
        col_widths(t, [CONTENT_W])
        c = t.rows[0].cells[0]
        shade(c, "F6F8FA")
        cell_borders(c, left=(18, ORANGE))
        lines = text.split("\n")
        first = True
        for line in lines:
            para = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            is_comment = line.strip().startswith(("#", "//"))
            run(para, line if line else " ", size=8, font=MONO, color=GREY if is_comment else INK)
        if len(lines) <= 14:
            keep_row_together(t.rows[0])
        self.spacer(6)

    def step(self, n: int, title: str, body: str, code: str | None = None, code_label: str | None = None):
        t = self.doc.add_table(rows=1, cols=2)
        table_borders(t, None)
        table_margins(t, 60, 60, 120, 120)
        col_widths(t, [1.1, CONTENT_W - 1.1])
        a, b = t.rows[0].cells
        shade(a, ORANGE)
        cell_text(a, str(n), bold=True, size=13, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(b, DARK)
        cell_text(b, title, bold=True, size=11, color="FFFFFF")
        keep_row_together(t.rows[0])
        for c in (a, b):
            c.paragraphs[0].paragraph_format.keep_with_next = True
        intro = self.p(body, after=4)
        intro.paragraph_format.space_before = Pt(4)
        intro.paragraph_format.keep_with_next = bool(code)
        if code:
            self.code(code, code_label)
        else:
            self.spacer(4)

    def data_table(self, headers, rows, widths, size=8.5, sev_col: int | None = None, zebra=True,
                   align_right: tuple = ()):
        t = self.doc.add_table(rows=1, cols=len(headers))
        table_borders(t, LINE)
        table_margins(t, 60, 60, 90, 90)
        col_widths(t, widths)
        hdr = t.rows[0]
        repeat_header(hdr)
        for c, h in zip(hdr.cells, headers):
            shade(c, DARK)
            cell_text(c, h, bold=True, size=size, color="FFFFFF")
        for i, r in enumerate(rows):
            cells = t.add_row().cells
            keep_row_together(t.rows[-1])
            for j, (c, v) in enumerate(zip(cells, r)):
                if zebra and i % 2 == 1:
                    shade(c, "FAFAFB")
                if sev_col is not None and j == sev_col:
                    shade(c, SEV_FILL.get(v, GREY))
                    cell_text(c, v, bold=True, size=size, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
                elif isinstance(v, tuple):  # (text, style dict)
                    cell_text(c, v[0], size=size, **v[1])
                else:
                    cell_text(c, v, size=size,
                              align=WD_ALIGN_PARAGRAPH.RIGHT if j in align_right else None)
        col_widths(t, widths)
        self.spacer(8)
        return t

    def kv_table(self, pairs, widths=(4.2, CONTENT_W - 4.2), keep=False):
        t = self.doc.add_table(rows=0, cols=2)
        table_borders(t, LINE)
        table_margins(t, 50, 50, 100, 100)
        for k, v in pairs:
            a, b = t.add_row().cells
            shade(a, LIGHT)
            cell_text(a, k, bold=True, size=9, color=DARK)
            if isinstance(v, tuple):
                cell_text(b, v[0], size=9, **v[1])
            else:
                cell_text(b, v, size=9)
            keep_row_together(t.rows[-1])
            if keep:
                for c in (a, b):
                    for para in c.paragraphs:
                        para.paragraph_format.keep_with_next = True
        col_widths(t, list(widths))
        if keep:
            self.spacer(6).paragraph_format.keep_with_next = True
        else:
            self.spacer(6)
        return t

    def kpis(self, items):
        t = self.doc.add_table(rows=2, cols=len(items))
        table_borders(t, None)
        table_margins(t, 120, 100, 120, 120)
        w = CONTENT_W / len(items)
        col_widths(t, [w] * len(items))
        for i, (value, label, color) in enumerate(items):
            top, bot = t.rows[0].cells[i], t.rows[1].cells[i]
            for c in (top, bot):
                shade(c, LIGHT)
                cell_borders(c, left=(4, "FFFFFF"), right=(4, "FFFFFF"))
            cell_borders(top, top=(24, color))
            cell_text(top, value, bold=True, size=24, color=color)
            cell_text(bot, label, size=8.5, color=GREY)
        self.spacer(10)


def add_toc(R: "Report", entries: list[str]):
    """TOC field whose cached result already lists the sections, so it reads well
    in Google Docs / LibreOffice / previews before Word recalculates page numbers."""
    def fld(p, kind, instr=None):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        if instr:
            r2 = p.add_run()
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            r2._r.append(it)

    paras = []
    for i, e in enumerate(entries):
        p = R.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        if i == 0:
            fld(p, "begin", ' TOC \\o "1-1" \\h \\z \\u ')
            fld(p, "separate")
        num, _, rest = e.partition(" ")
        run(p, num, bold=True, size=11.5, color=ORANGE)
        run(p, "\t" + rest, bold=True, size=11.5, color=DARK)
        para_border(p, "bottom", 2, LINE, 6)
        paras.append(p)
    fld(paras[-1], "end")


# ---- chart -----------------------------------------------------------------

def stream_risk_chart(findings: list[dict]) -> io.BytesIO:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    per = defaultdict(lambda: Counter())
    for f in findings:
        per[f["stream"]][f["severity"]] += 1
    streams = sorted(per, key=lambda s: (-(per[s]["Critical"] * 100 + per[s]["High"] * 10 + per[s]["Medium"]), s))
    labels = [s.split('app="')[-1].rstrip('"}') if 'app="' in s else s for s in streams]

    fig, ax = plt.subplots(figsize=(7.2, 0.42 * len(streams) + 0.9), dpi=200)
    left = [0] * len(streams)
    for sev in ("Critical", "High", "Medium", "Low"):
        vals = [per[s][sev] for s in streams]
        ax.barh(labels, vals, left=left, color="#" + SEV_FILL[sev], label=sev, height=0.6, edgecolor="white")
        left = [a + b for a, b in zip(left, vals)]
    ax.invert_yaxis()
    ax.set_xlabel("Findings (one per data type per stream)", fontsize=8, color="#" + GREY)
    ax.tick_params(labelsize=8, colors="#" + INK)
    ax.xaxis.get_major_locator().set_params(integer=True)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("#" + LINE)
    ax.grid(axis="x", color="#" + LINE, linewidth=0.6)
    ax.set_axisbelow(True)
    ax.legend(ncol=4, fontsize=7.5, frameon=False, loc="lower center", bbox_to_anchor=(0.5, 1.0))
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def short(stream: str) -> str:
    return stream.replace('", ', '",\u200b')


def fmt_int(n) -> str:
    return f"{int(n):,}"


def pct(x: float) -> str:
    return f"{x * 100:.1f}%" if x >= 0.001 else "<0.1%"


# ---- the report ------------------------------------------------------------

def render(data: dict, out_path: str):
    R = Report(data)
    meta, summ, F = data["meta"], data["summary"], data["findings"]
    gen = datetime.fromisoformat(meta["generated_at"])

    # ---------- cover
    cover = R.doc.add_table(rows=1, cols=1)
    table_borders(cover, None)
    table_margins(cover, 500, 500, 500, 500)
    col_widths(cover, [CONTENT_W])
    c = cover.rows[0].cells[0]
    shade(c, DARK)
    cell_borders(c, top=(48, ORANGE), bottom=(48, YELLOW))
    p0 = c.paragraphs[0]
    run(p0, "Grafana Labs", bold=True, size=14, color="FFFFFF")
    run(p0, ".", bold=True, size=14, color=ORANGE)
    for txt, kw in (
        ("", {}),
        (meta.get("title", "Telemetry Privacy Report"), dict(bold=True, size=30, color="FFFFFF")),
        (meta.get("subtitle", "Personal data & secrets discovered in log pipelines — DPIA annex"),
         dict(size=13, color=YELLOW)),
        ("", {}),
        ("Privacy Audit  │  GDPR Data Map  │  Remediation Runbook", dict(bold=True, size=10, color="FFFFFF")),
        ("Loki · LogQL · Grafana Alloy · OpenTelemetry Collector", dict(size=10, color="AEB4BB")),
        ("", {}),
        (f"{meta.get('organization', '')}  ·  {meta.get('environment', '')}  ·  Version {meta.get('version', '1.0')}  ·  "
         f"{gen.strftime('%B %Y')}", dict(size=10, color="AEB4BB")),
    ):
        para = c.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        if txt:
            run(para, txt, **kw)

    R.spacer(18)
    R.kv_table([
        ("Scope", meta.get("scope", "")),
        ("Observation window", f"{meta['window_start'][:16].replace('T', ' ')} → {meta['window_end'][:16].replace('T', ' ')} UTC"),
        ("Data source", meta.get("source", "")),
        ("Prepared by", meta.get("author", "")),
        ("Generated", gen.strftime("%Y-%m-%d %H:%M UTC") + f"  ·  grafana-privacy-report v{meta.get('tool_version', '')}"),
        ("Distribution", (meta.get("distribution", "DPO, Security, Platform Engineering — do not forward"),
                          dict(bold=True, color=SEV_FILL["Critical"]))),
    ])
    R.callout("INFO", "All values in this report are masked at scan time. Raw personal data never leaves the "
                      "scanner process; repeated values are counted via one-way SHA-256 fingerprints.",
              "Privacy by design.")
    R.page_break()

    # ---------- TOC
    h = R.doc.add_heading(level=1)
    run(h, "Table of Contents")
    add_toc(R, ["1. Executive Summary", "2. Scope and Methodology", "3. Findings Register",
                "4. Detailed Findings", "5. Remediation Plan", "6. Data Inventory and Legal Basis",
                "7. Review and Sign-off", "8. Troubleshooting"])
    R.callout("TIP", "Word refreshes this table with page numbers when the document opens (accept the prompt), "
                     "or right-click it and choose Update Field.", "Page numbers.")
    R.page_break()

    # ---------- 1. Executive summary
    sev = summ["by_severity"]
    R.h1("1. Executive Summary")
    R.p(f"This report documents personal data and credentials found inside the log pipeline of "
        f"{meta.get('organization', 'the organization')}. {fmt_int(summ['lines_sampled'])} log lines were sampled "
        f"from {summ['streams_scanned']} Loki streams. {summ['streams_affected']} streams contain at least one category "
        f"of personal data or secrets, producing {summ['findings']} findings. Based on sampled hit-rates, an estimated "
        f"{fmt_int(summ['est_pii_lines_per_day'])} personal-data detections per day flow into long-term storage, where "
        f"they become queryable by everyone with access to the Loki data source.")
    R.kpis([
        (str(summ["streams_scanned"]), "Streams scanned", DARK),
        (str(summ["streams_affected"]), "Streams with personal data", ORANGE),
        (str(sev["Critical"]), "Critical findings", SEV_FILL["Critical"]),
        (str(sev["High"]), "High findings", SEV_FILL["High"]),
    ])
    R.h2("1.1  Findings by Stream and Severity")
    if F:
        R.doc.add_picture(stream_risk_chart(F), width=Cm(16))
        R.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        R.callout("TIP", "None of the enabled detectors matched any sampled line. Keep this report as evidence "
                         "and schedule the scan regularly so new leaks are caught early.", "Clean result.")

    R.h2("1.2  Key Messages")
    creds = [f for f in F if BY_KEY[f["detector"]].category == "Credential"]
    cards = [f for f in F if f["detector"] == "card"]
    long_ret = [f for f in F if f["retention_days"] > 30]
    if creds:
        R.callout("WARN", f"{len(creds)} finding(s) expose live credentials ({', '.join(sorted({f['label'] for f in creds}))}). "
                          "Treat these as security incidents: rotate the credentials before purging the logs.",
                  "Credentials in logs.")
    if cards:
        R.callout("WARN", "Payment card numbers were found in "
                          f"{', '.join(sorted({f['stream'] for f in cards}))}. Any system storing full PANs is in "
                          "PCI DSS scope; this includes the Loki object store and its backups.",
                  "PCI DSS scope creep.")
    if long_ret:
        R.callout("INFO", f"{len({f['stream'] for f in long_ret})} affected stream(s) keep data longer than 30 days. "
                          "Shorter retention for these streams lowers risk immediately, even before redaction ships.",
                  "Retention amplifies exposure.")
    if F:
        R.callout("TIP", "Section 5 contains generated, copy-paste configuration for Grafana Alloy and the OpenTelemetry "
                     "Collector that redacts exactly the patterns found, scoped to the affected streams only.",
                  "Fast path.")

    # ---------- 2. Scope & methodology
    R.page_break()
    R.h1("2. Scope and Methodology")
    R.p("The scanner reads log lines through the Loki HTTP API using read-only queries, applies pattern detectors "
        "with validation, and aggregates results per stream and data type. Nothing is written back to Loki.")
    R.h2("2.1  Scan Parameters")
    R.kv_table([
        ("Loki endpoint", meta.get("source", "")),
        ("Stream selector(s)", ", ".join(meta.get("selectors", []))),
        ("Grouped by labels", ", ".join(meta.get("group_by", []))),
        ("Window", meta.get("window", "")),
        ("Sample size", f"Up to {fmt_int(meta.get('sample_lines', 0))} most recent lines per stream"),
        ("Detectors", ", ".join(BY_KEY[k].label for k in meta.get("detectors", BY_KEY.keys()))),
    ])
    R.h2("2.2  Detector Catalog")
    R.p("Each detector combines a regular expression with a validator where one exists, which removes most false "
        "positives (for example, random 16-digit order IDs fail the Luhn check and are not reported as cards).",
        after=6)
    validators = {"card": "Luhn checksum", "iban": "ISO 13616 mod-97", "ip": "Public ranges only",
                  "jwt": "3-part base64url", "aws_key": "AKIA/ASIA prefix", "bearer": "Header context"}
    R.data_table(["Data type", "Category", "Validation", "Weight", "Legal / risk context"],
                 [[d.label, d.category, validators.get(d.key, "Pattern only"), str(d.weight), d.gdpr_note]
                  for d in BY_KEY.values() if d.key in meta.get("detectors", BY_KEY.keys())],
                 [3.3, 3.1, 2.7, 1.4, 6.5], align_right=(3,))
    R.h2("2.3  Risk Model")
    R.p("Risk score = detector weight × volume factor × retention factor. The model is deliberately simple so that "
        "auditors can recompute any score by hand.", after=6)
    R.data_table(["Factor", "Values", "Rationale"], [
        ["Detector weight", "1 – 5 (see 2.2)", "Sensitivity of the data type if exposed."],
        ["Volume factor", "1 (<100/day) · 2 (<10k/day) · 3 (≥10k/day)", "More records means more data subjects affected."],
        ["Retention factor", "1.0 (≤7 d) · 1.5 (≤30 d) · 2.0 (>30 d)", "Longer storage widens the exposure window."],
        ["Severity bands", "Critical ≥20 · High ≥12 · Medium ≥6 · Low <6", "Maximum possible score is 30."],
    ], [3.3, 6.4, 7.3])
    R.callout("WARN", "Pattern detection cannot find free-text personal data such as names, postal addresses or "
                      "health details, and it misses values that are base64-encoded or split across lines. "
                      "A clean result is evidence of absence for the listed patterns only, not proof that a stream "
                      "is free of personal data.", "Limitations.")

    # ---------- 3. Findings register
    R.h1("3. Findings Register")
    R.p("One row per data type per stream, ordered by severity. Hit rate is the share of sampled lines containing "
        "the data type; the daily estimate extrapolates that rate to the stream's total volume.")
    R.data_table(["ID", "Stream", "Data type", "Hit rate", "Est. lines / day", "Retention", "Severity"],
                 [[f["id"], short(f["stream"]), f["label"], pct(f["hit_rate"]), fmt_int(f["est_lines_per_day"]),
                   f"{f['retention_days']:g} d", f["severity"]] for f in F],
                 [1.6, 5.0, 3.0, 1.6, 2.1, 1.7, 2.0], size=8, sev_col=6, align_right=(3, 4, 5))

    # ---------- 4. Detailed findings
    if not F:
        R.page_break()
        R.h1("4. Detailed Findings")
        R.p("No findings to detail.")
        R.h1("5. Remediation Plan")
        R.p("No remediation required for the patterns and streams in scope.")
        creds = []
    else:
        _detail_and_fix(R, F, meta, creds)
    _inventory_signoff_troubleshooting(R, F)
    R.doc.save(out_path)
    return out_path


def _detail_and_fix(R, F, meta, creds):
    R.page_break()
    R.h1("4. Detailed Findings")
    detailed = [f for f in F if f["severity"] in ("Critical", "High")]
    R.p(f"Critical and High findings ({len(detailed)}) are documented individually below. Medium and Low findings are "
        "covered by the register in Section 3 and by the stream-level fixes in Section 5.")
    for f in detailed:
        R.h2(f"{f['id']}  ·  {f['label']} in {f['stream']}")
        R.kv_table([
            ("Severity", (f"{f['severity']}  (risk score {f['risk_score']:g} / 30)",
                          dict(bold=True, color=SEV_FILL[f["severity"]]))),
            ("Stream owner", f["owner"] or "[to be assigned]"),
            ("Category", f"{f['category']} — {f['gdpr_note']}"),
            ("Evidence", f"{fmt_int(f['lines_with_hit'])} of {fmt_int(f['sampled_lines'])} sampled lines "
                         f"({pct(f['hit_rate'])}), {fmt_int(f['distinct_values'])} distinct value(s)"),
            ("Masked examples", ("   ".join(f["examples"]), dict(font=MONO))),
            ("Exposure", f"≈ {fmt_int(f['est_lines_per_day'])} lines/day, retained {f['retention_days']:g} days"),
        ], keep=True)
        R.code(rem.logql_check(f), "Verify — count remaining matches (should reach 0 after the fix):")

    # ---------- 5. Remediation plan
    R.page_break()
    R.h1("5. Remediation Plan")
    R.p("Apply the steps in order. Steps 2 and 5 reduce risk within hours; step 1 removes the root cause and should "
        "be tracked as engineering work for each stream owner.")
    owners = sorted({(f["owner"] or "Unassigned") for f in F})
    R.step(1, "Stop the leak at the source",
           "Ask each stream owner to remove personal data from log statements (log IDs instead of emails, never log "
           "request bodies or auth headers). Owners with open findings: " + ", ".join(owners) + ".")
    R.step(2, "Redact in the pipeline with Grafana Alloy",
           "Add this loki.process component between your log sources and loki.write. Each stage.match is scoped "
           "to one affected stream, so unaffected streams keep full fidelity (private IPs used for debugging are "
           "not touched elsewhere).", rem.alloy_config(F), "config.alloy")
    R.callout("TIP", "Using the OpenTelemetry Collector instead? The equivalent transform processor is generated "
                     "below. It applies to all logs in the pipeline; add a where clause per statement to scope it.",
              "OTel alternative.")
    R.code(rem.otel_config(F), "otel-collector.yaml")
    sample = detailed[0] if detailed else (F[0] if F else None)
    if sample:
        R.step(3, "Verify that redaction works",
               "Run the LogQL check from Section 4 for each finding, in Grafana Explore or as a recording rule. "
               "Turn it into an alert so regressions are caught when new code ships:",
               f"# alert: PIILeakDetected — fires if redaction regresses\n{rem.logql_check(sample)} > 0")
        R.step(4, "Purge data already stored",
               "Redaction only protects new data. Use Loki's delete API to remove historical lines that match. "
               "Requires the compactor with retention enabled and deletion_mode set to filter-and-delete. Example "
               "for " + sample["id"] + ":",
               rem.delete_request(sample, meta["window_start"][:19] + "Z", meta["window_end"][:19] + "Z"))
    R.step(5, "Shorten retention for high-risk streams",
           "Per-stream retention limits exposure while fixes roll out. Generated for streams with Critical or High "
           "findings (7 days shown; agree the period with the DPO):", rem.retention_config(F, 7), "loki.yaml")
    if creds:
        R.step(6, "Rotate exposed credentials",
               "Credentials that reached the log store must be assumed compromised: revoke and reissue them, "
               "invalidate active sessions for leaked JWTs, and review access logs for misuse during the window.")
    R.callout("WARN", "Purging logs that are subject to a legal hold or an open incident investigation may be "
                      "prohibited. Confirm with Legal before running delete requests.", "Before you delete.")



def _inventory_signoff_troubleshooting(R, F):
    R.page_break()
    R.h1("6. Data Inventory and Legal Basis")
    R.p("This section is the DPIA annex. The scanner pre-fills what it observed; the Data Protection Officer "
        "completes purpose, lawful basis and the decision for each stream.")
    by_stream = defaultdict(list)
    for f in F:
        by_stream[f["stream"]].append(f)
    todo = ("[DPO to complete]", dict(italic=True, color=GREY))
    R.data_table(["Stream", "Owner", "Personal data observed", "Retention", "Purpose", "Lawful basis", "Decision"],
                 [[short(s), fs[0]["owner"], ", ".join(sorted({x["label"] for x in fs})), f"{fs[0]['retention_days']:g} d",
                   todo, todo, todo] for s, fs in by_stream.items()],
                 [3.6, 2.0, 3.4, 1.5, 2.2, 2.2, 2.1], size=8)

    # ---------- 7. Sign-off
    R.h1("7. Review and Sign-off")
    R.data_table(["Role", "Name", "Decision", "Signature", "Date"],
                 [[r, "", "", "", ""] for r in ("Data Protection Officer", "Security Lead",
                                                  "Platform / Observability Owner", "Engineering Manager")],
                 [4.6, 3.6, 3.2, 3.2, 2.4], size=9, zebra=False)

    # ---------- 8. Troubleshooting
    R.page_break()
    R.h1("8. Troubleshooting")
    R.p("Common issues when running the scanner and applying the fixes.")
    R.data_table(["Error / Symptom", "Cause", "Fix"], [
        ["401 / 403 from Loki", "Missing tenant header or token without logs:read scope.",
         "Set loki.tenant (X-Scope-OrgID); on Grafana Cloud use instance ID + access-policy token."],
        ["Sample smaller than requested", "Query limit capped by max_entries_limit_per_query (default 5000).",
         "Lower sample_lines_per_stream or raise the limit for the scanner's tenant."],
        ["Card findings on order IDs", "A numeric ID happens to pass the Luhn check.",
         "Add the stream/field to allowlist in config.yaml; risk is low but review once."],
        ["Alloy replace has no effect", "Stage placed after loki.write, or stream selector does not match.",
         "Check component wiring in the Alloy UI; test the selector in Explore first."],
        ["Delete request stays pending", "Compactor retention or deletion_mode not enabled.",
         "Enable compactor.retention_enabled and set deletion_mode: filter-and-delete."],
        ["Redaction breaks a dashboard", "A panel or alert parsed the value that is now redacted.",
         "Log a stable pseudonymous ID (e.g. user_id) and query that instead."],
    ], [4.0, 6.0, 7.0])
    R.p("This report supports, but does not replace, a formal Data Protection Impact Assessment. It is not legal "
        "advice. Grafana Labs and the Grafana logo are trademarks of Grafana Labs; this community template is not "
        "an official Grafana Labs publication.", italic=True, size=8, color=GREY)
